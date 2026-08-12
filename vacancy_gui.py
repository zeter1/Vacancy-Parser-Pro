import json
import os
import sys
import threading
import traceback
import time
import urllib.parse
import tkinter as tk
import webbrowser
from datetime import datetime
from pathlib import Path
from statistics import median
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from job_scraper import JobScraper
from problem_logging import SearchProblemLogger


APP_VERSION = "4.9"
APP_NAME = "Vacancy Parser Pro"
BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "parser_config.json"
PROFILES_FILE = BASE_DIR / "saved_params.json"
ICON_FILE = BASE_DIR / "icon.ico"
LOG_DIR = BASE_DIR / "Логи проблем"


class App:
    def __init__(self, root):
        self.root = root
        self.root.title(f"Парсер вакансий v{APP_VERSION} PRO | by Дмитрий Колесниченко")
        self.root.geometry("1300x850")
        self.root.minsize(1000, 700)

        try:
            if ICON_FILE.exists():
                self.root.iconbitmap(str(ICON_FILE))
        except tk.TclError:
            pass

        self.sort_reverse = False
        self.last_sort_col = None
        self.text_widgets = []
        self.found_data = []
        self.search_log = []
        self.last_search_error = None
        self.last_search_traceback = None
        self.search_cancel_event = threading.Event()
        self.search_thread = None
        self.current_log_session_dir = None
        self._vacancy_open_in_progress = False
        self.problem_logger = SearchProblemLogger(
            LOG_DIR, APP_NAME, APP_VERSION, keep_sessions=15
        )

        self.config = self.load_config()
        self.saved_params = self.load_saved_params()
        self.scraper = JobScraper(log_callback=self._append_search_log)

        self.setup_theme()
        self.create_widgets()
        self.restore_last_settings()

        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)

    def setup_theme(self):
        """Настройка темы оформления"""
        style = ttk.Style()
        style.theme_use('clam')
        
        # Настройка цветов
        bg_color = '#f0f0f0'
        fg_color = '#000000'
        select_bg = '#0078d7'
        select_fg = '#ffffff'
        
        style.configure("Treeview", 
                       rowheight=28, 
                       font=('Segoe UI', 10),
                       background='white',
                       foreground=fg_color,
                       fieldbackground='white')
        style.configure("Treeview.Heading", 
                       font=('Segoe UI', 10, 'bold'),
                       background='#e0e0e0',
                       foreground=fg_color)
        style.map('Treeview', 
                 background=[('selected', select_bg)],
                 foreground=[('selected', select_fg)])
        
        style.configure("TButton", 
                       font=('Segoe UI', 10),
                       padding=6)
        style.configure("TLabel", 
                       font=('Segoe UI', 10))
        style.configure("TLabelframe.Label", 
                       font=('Segoe UI', 11, 'bold'))

    def create_widgets(self):
        """Создание всех виджетов интерфейса"""
        # Верхняя панель с профилями
        self.create_profile_frame()
        
        # Панель выбора сайтов
        self.create_sites_frame()
        
        # Критерии поиска
        self.create_search_criteria_frame()
        
        # Строка состояния и прогресс-бар
        self.create_status_frame()
        
        # Таблица результатов
        self.create_results_frame()
        
        # Нижняя панель с кнопками
        self.create_buttons_frame()
        
        # Контекстные меню
        self.create_context_menus()

    def create_profile_frame(self):
        """Фрейм для сохраненных параметров"""
        frame = tk.LabelFrame(self.root, text="📋 Профили поиска", padx=15, pady=10)
        frame.pack(fill=tk.X, padx=10, pady=5)
        
        tk.Label(frame, text="Название:").grid(row=0, column=0, sticky="w", pady=5)
        self.entry_profile_name = ttk.Entry(frame, width=20)
        self.entry_profile_name.grid(row=0, column=1, sticky="w", padx=5)
        self.text_widgets.append(self.entry_profile_name)
        
        self.btn_save_profile = ttk.Button(frame, text="💾 Сохранить", command=self.save_search_params)
        self.btn_save_profile.grid(row=0, column=2, padx=10)
        
        tk.Label(frame, text="Загрузить:").grid(row=0, column=3, sticky="w", padx=10)
        self.cb_profiles = ttk.Combobox(frame, values=list(self.saved_params.keys()), width=20, state="readonly")
        self.cb_profiles.grid(row=0, column=4, sticky="w", padx=5)
        self.cb_profiles.bind('<<ComboboxSelected>>', self.load_search_params)
        
        self.btn_delete_profile = ttk.Button(frame, text="🗑️ Удалить", command=self.delete_search_params)
        self.btn_delete_profile.grid(row=0, column=5, padx=10)

    def create_sites_frame(self):
        """Выбор реально поддерживаемых источников."""
        frame = tk.LabelFrame(self.root, text="🌐 Источники поиска", padx=15, pady=10)
        frame.pack(fill=tk.X, padx=10, pady=5)

        tk.Label(
            frame,
            text="Показываются только источники, для которых в программе действительно реализован поиск.",
            fg="#666666",
            font=("Segoe UI", 8),
        ).pack(anchor=tk.W, pady=(0, 5))

        self.site_vars = {}

        belarus_frame = tk.Frame(frame)
        belarus_frame.pack(fill=tk.X, pady=3)
        tk.Label(belarus_frame, text="🇧🇾 Беларусь:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(0, 10))
        for site_name in ("Rabota.by", "Praca.by", "Belmeta", "GSZ.gov.by"):
            var = tk.BooleanVar(value=True)
            self.site_vars[site_name] = var
            tk.Checkbutton(belarus_frame, text=site_name, variable=var, font=("Segoe UI", 9)).pack(side=tk.LEFT, padx=5)

        russia_frame = tk.Frame(frame)
        russia_frame.pack(fill=tk.X, pady=3)
        tk.Label(russia_frame, text="🇷🇺 Россия:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(0, 10))
        var = tk.BooleanVar(value=True)
        self.site_vars["HH.ru"] = var
        tk.Checkbutton(russia_frame, text="HH.ru", variable=var, font=("Segoe UI", 9)).pack(side=tk.LEFT, padx=5)

        buttons_frame = tk.Frame(frame)
        buttons_frame.pack(fill=tk.X, pady=(5, 0))
        ttk.Button(buttons_frame, text="✓ Все", command=self.select_all_sites, width=10).pack(side=tk.LEFT, padx=5)
        ttk.Button(buttons_frame, text="✗ Снять", command=self.deselect_all_sites, width=10).pack(side=tk.LEFT, padx=5)
        ttk.Button(buttons_frame, text="🇧🇾 Беларусь", command=self.select_belarus_sites, width=12).pack(side=tk.LEFT, padx=5)
        ttk.Button(buttons_frame, text="🇷🇺 Россия", command=self.select_russia_sites, width=12).pack(side=tk.LEFT, padx=5)

    def create_search_criteria_frame(self):
        """Фрейм критериев поиска."""
        frame = tk.LabelFrame(self.root, text="🔍 Критерии поиска", padx=15, pady=10)
        frame.pack(fill=tk.X, padx=10, pady=5)

        tk.Label(frame, text="Профессии:").grid(row=0, column=0, sticky="w", pady=5)
        self.entry_prof = ttk.Entry(frame, width=40)
        self.entry_prof.grid(row=0, column=1, sticky="ew", padx=5)
        self.entry_prof.insert(0, "Python")
        self.text_widgets.append(self.entry_prof)

        tk.Label(frame, text="Города:").grid(row=0, column=2, sticky="w", padx=10)
        self.entry_city = ttk.Entry(frame, width=30)
        self.entry_city.grid(row=0, column=3, sticky="ew", padx=5)
        self.entry_city.insert(0, "Минск")
        self.text_widgets.append(self.entry_city)

        tk.Label(frame, text="Мин. доход:").grid(row=1, column=0, sticky="w", pady=5)
        self.entry_salary = ttk.Entry(frame, width=15)
        self.entry_salary.grid(row=1, column=1, sticky="w", padx=5)
        self.text_widgets.append(self.entry_salary)

        tk.Label(frame, text="Валюта:").grid(row=1, column=2, sticky="w", padx=5)
        self.cb_currency = ttk.Combobox(frame, values=["BYN", "USD", "RUB", "EUR"], state="readonly", width=10)
        self.cb_currency.set("BYN")
        self.cb_currency.grid(row=1, column=3, sticky="w", padx=5)

        tk.Label(frame, text="Опыт:").grid(row=2, column=0, sticky="w", padx=10)
        self.cb_exp = ttk.Combobox(
            frame,
            values=["Не имеет значения", "Нет опыта", "От 1 года до 3 лет", "От 3 до 6 лет", "Более 6 лет"],
            state="readonly",
            width=28,
        )
        self.cb_exp.current(0)
        self.cb_exp.grid(row=2, column=1, sticky="w", padx=5)

        tk.Label(frame, text="Формат / график:").grid(row=2, column=2, sticky="w", pady=5)
        self.cb_sched = ttk.Combobox(
            frame,
            values=[
                "Любой",
                "Удаленная работа",
                "На месте работодателя",
                "Гибрид",
                "Полный день",
                "Сменный график",
                "Гибкий график",
            ],
            state="readonly",
            width=25,
        )
        self.cb_sched.current(0)
        self.cb_sched.grid(row=2, column=3, sticky="w", padx=5)

        tk.Label(frame, text="Период:").grid(row=3, column=0, sticky="w", padx=10)
        self.cb_period = ttk.Combobox(
            frame,
            values=["За 30 дней", "За неделю", "За 3 дня", "За сутки"],
            state="readonly",
            width=28,
        )
        self.cb_period.set("За неделю")
        self.cb_period.grid(row=3, column=1, sticky="w", padx=5)

        tk.Label(frame, text="Исключить:").grid(row=3, column=2, sticky="w", pady=5)
        self.entry_exclude = ttk.Entry(frame, width=30)
        self.entry_exclude.grid(row=3, column=3, sticky="ew", padx=5)
        self.text_widgets.append(self.entry_exclude)

        tk.Label(frame, text="Фильтр зарплаты:").grid(row=4, column=0, sticky="w", pady=5)
        self.cb_salary_mode = ttk.Combobox(
            frame,
            values=["Не фильтровать", "Мягкий", "Строгий"],
            state="readonly",
            width=28,
        )
        self.cb_salary_mode.set("Мягкий")
        self.cb_salary_mode.grid(row=4, column=1, sticky="w", padx=5)
        tk.Label(
            frame,
            text=(
                "Мягкий: вакансии без зарплаты остаются; диапазон может доходить до минимума. "
                "Строгий: указанная нижняя граница должна быть не ниже минимума."
            ),
            fg="#666666",
            font=("Segoe UI", 8),
            wraplength=520,
            justify=tk.LEFT,
        ).grid(row=4, column=2, columnspan=2, sticky="w", padx=5)

        action_frame = tk.Frame(frame)
        action_frame.grid(row=0, column=4, rowspan=5, padx=20, sticky="ns")
        self.btn_search = ttk.Button(action_frame, text="🔍 НАЙТИ", command=self.start_search_thread, width=18)
        self.btn_search.pack(fill=tk.X, pady=(2, 8))
        self.btn_stop = ttk.Button(action_frame, text="■ ОСТАНОВИТЬ", command=self.cancel_search, width=18)
        self.btn_stop.pack(fill=tk.X)
        self.btn_stop.config(state=tk.DISABLED)

        frame.grid_columnconfigure(1, weight=1)
        frame.grid_columnconfigure(3, weight=1)

    def create_status_frame(self):
        """Фрейм состояния и прогресс-бара"""
        frame = tk.Frame(self.root)
        frame.pack(fill=tk.X, padx=10, pady=5)
        
        self.lbl_status = tk.Label(frame, text="Готов к поиску", fg="gray", font=('Segoe UI', 9))
        self.lbl_status.pack(anchor=tk.W)
        
        self.progress = ttk.Progressbar(frame, orient='horizontal', mode='determinate', length=400)
        self.progress.pack(fill=tk.X, pady=5)
        self.progress['value'] = 0

    def create_results_frame(self):
        """Фрейм таблицы результатов"""
        frame = tk.Frame(self.root)
        frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        cols = ("position", "salary", "company", "city", "source", "date")
        self.tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode='extended')
        
        self.tree.heading("position", text="📋 Должность", command=lambda: self.sort_column("position", False))
        self.tree.heading("salary", text="💰 Зарплата", command=lambda: self.sort_column("salary", False))
        self.tree.heading("company", text="🏢 Компания", command=lambda: self.sort_column("company", False))
        self.tree.heading("city", text="🌍 Город", command=lambda: self.sort_column("city", False))
        self.tree.heading("source", text="🌐 Источник", command=lambda: self.sort_column("source", False))
        self.tree.heading("date", text="📅 Дата", command=lambda: self.sort_column("date", False))
        
        self.tree.column("position", width=280)
        self.tree.column("salary", width=190, minwidth=165)
        self.tree.column("company", width=200)
        self.tree.column("city", width=110)
        self.tree.column("source", width=130)
        self.tree.column("date", width=90)

        scrollbar_y = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscroll=scrollbar_y.set, xscroll=scrollbar_x.set)
        
        self.tree.grid(row=0, column=0, sticky='nsew')
        scrollbar_y.grid(row=0, column=1, sticky='ns')
        scrollbar_x.grid(row=1, column=0, sticky='ew')
        
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        
        self.tree.bind("<Double-1>", self.on_double_click)
        self.tree.bind("<Button-3>", self.show_context_menu)
        self.tree.bind("<Control-c>", self.copy_selected_to_clipboard)

    def create_buttons_frame(self):
        """Нижняя панель действий."""
        frame = tk.Frame(self.root)
        frame.pack(pady=10)

        ttk.Button(frame, text="💾 Сохранить в Excel", command=self.save_to_excel, width=20).pack(side=tk.LEFT, padx=8)
        self.btn_open = ttk.Button(frame, text="🔗 Открыть вакансию", command=self.open_selected_vacancy, width=20)
        self.btn_open.pack(side=tk.LEFT, padx=8)
        self.btn_open.config(state=tk.DISABLED)

        ttk.Button(frame, text="📊 Статистика", command=self.show_statistics, width=15).pack(side=tk.LEFT, padx=8)
        ttk.Button(frame, text="🔄 Перерисовать", command=self.refresh_results, width=15).pack(side=tk.LEFT, padx=8)
        ttk.Button(frame, text="📝 Лог поиска", command=self.show_search_log, width=15).pack(side=tk.LEFT, padx=8)

    def create_context_menus(self):
        """Создание контекстных меню"""
        # Меню для таблицы
        self.context_menu = tk.Menu(self.root, tearoff=0)
        self.context_menu.add_command(label="🔗 Открыть вакансию", command=self.open_selected_vacancy)
        self.context_menu.add_command(label="📋 Копировать ссылку", command=self.copy_link)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="📊 Информация", command=self.show_vacancy_info)

        # Меню для текстовых полей
        self.text_context_menu = tk.Menu(self.root, tearoff=0)
        self.text_context_menu.add_command(label="Копировать", command=self.copy_text)
        self.text_context_menu.add_command(label="Вставить", command=self.paste_text)
        self.text_context_menu.add_separator()
        self.text_context_menu.add_command(label="Выделить все", command=self.select_all_text)

        # Привязка контекстного меню к текстовым полям
        for widget in self.text_widgets:
            widget.bind("<Button-3>", self.show_text_context_menu)
            widget.bind("<Control-Key>", self.on_key_press)

    # --- МЕТОДЫ РАБОТЫ С КОНФИГУРАЦИЕЙ ---
    
    def load_config(self):
        """Безопасная загрузка конфигурации."""
        try:
            if CONFIG_FILE.exists():
                with CONFIG_FILE.open("r", encoding="utf-8") as f:
                    data = json.load(f)
                return data if isinstance(data, dict) else {}
        except (OSError, json.JSONDecodeError) as exc:
            print(f"Ошибка загрузки конфигурации: {exc}")
        return {}

    def save_config(self):
        """Атомарное сохранение конфигурации, чтобы не оставить битый JSON."""
        try:
            CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
            temp_file = CONFIG_FILE.with_suffix(CONFIG_FILE.suffix + ".tmp")
            with temp_file.open("w", encoding="utf-8") as f:
                json.dump(self.config, f, ensure_ascii=False, indent=2)
            os.replace(temp_file, CONFIG_FILE)
        except OSError as exc:
            print(f"Ошибка сохранения конфигурации: {exc}")

    def load_saved_params(self):
        """Загрузка сохранённых профилей поиска."""
        try:
            if PROFILES_FILE.exists():
                with PROFILES_FILE.open("r", encoding="utf-8") as f:
                    data = json.load(f)
                return data if isinstance(data, dict) else {}
        except (OSError, json.JSONDecodeError) as exc:
            print(f"Ошибка загрузки профилей: {exc}")
        return {}

    def save_saved_params(self):
        """Атомарное сохранение профилей поиска."""
        try:
            PROFILES_FILE.parent.mkdir(parents=True, exist_ok=True)
            temp_file = PROFILES_FILE.with_suffix(PROFILES_FILE.suffix + ".tmp")
            with temp_file.open("w", encoding="utf-8") as f:
                json.dump(self.saved_params, f, ensure_ascii=False, indent=2)
            os.replace(temp_file, PROFILES_FILE)
        except OSError as exc:
            print(f"Ошибка сохранения профилей: {exc}")

    def save_current_settings(self):
        """Сохранение текущих настроек."""
        salary_mode = self.cb_salary_mode.get() or "Мягкий"
        self.config["last_settings"] = {
            "queries": self.entry_prof.get().strip(),
            "city": self.entry_city.get().strip(),
            "salary": self.entry_salary.get().strip(),
            "salary_currency": self.cb_currency.get(),
            "experience": self.cb_exp.get(),
            "schedule": self.cb_sched.get(),
            "period": self.cb_period.get(),
            "exclude": self.entry_exclude.get().strip(),
            "enabled_sites": self.get_enabled_sites(),
            "salary_mode": salary_mode,
            # Compatibility for older copies/config readers.
            "strict_salary": salary_mode == "Строгий",
        }
        self.save_config()

    def restore_last_settings(self):
        """Восстановление последних настроек с миграцией старых значений."""
        settings = self.config.get("last_settings")
        if not isinstance(settings, dict):
            return

        self.entry_prof.delete(0, tk.END)
        self.entry_prof.insert(0, settings.get("queries", "Python"))

        self.entry_city.delete(0, tk.END)
        self.entry_city.insert(0, settings.get("city", "Минск"))

        self.entry_salary.delete(0, tk.END)
        self.entry_salary.insert(0, settings.get("salary", ""))

        self.entry_exclude.delete(0, tk.END)
        self.entry_exclude.insert(0, settings.get("exclude", ""))

        self.cb_currency.set(settings.get("salary_currency", "BYN"))
        self.cb_exp.set(settings.get("experience", "Не имеет значения"))

        schedule = settings.get("schedule", "Любой")
        if schedule not in self.cb_sched["values"]:
            schedule = "Любой"
        self.cb_sched.set(schedule)

        period = settings.get("period", "За неделю")
        if period in {"За все время", "За месяц"}:
            period = "За 30 дней"
        if period not in self.cb_period["values"]:
            period = "За неделю"
        self.cb_period.set(period)

        salary_mode = settings.get("salary_mode")
        if salary_mode not in self.cb_salary_mode["values"]:
            salary_mode = "Строгий" if settings.get("strict_salary") else "Мягкий"
        self.cb_salary_mode.set(salary_mode)

        enabled_sites = settings.get("enabled_sites")
        if not isinstance(enabled_sites, list):
            enabled_sites = list(self.site_vars)
        for site_name, var in self.site_vars.items():
            var.set(site_name in enabled_sites)

    def on_closing(self):
        """Корректное завершение приложения."""
        self.search_cancel_event.set()
        try:
            self.save_current_settings()
        finally:
            try:
                self.scraper.close()
            except Exception:
                pass
            self.root.destroy()

    # --- МЕТОДЫ РАБОТЫ С ПРОФИЛЯМИ ---
    
    def save_search_params(self):
        """Сохранение профиля"""
        profile_name = self.entry_profile_name.get().strip()
        if not profile_name:
            messagebox.showwarning("Внимание", "Введите название профиля!")
            return

        salary_mode = self.cb_salary_mode.get() or "Мягкий"
        params = {
            "queries": self.entry_prof.get(),
            "city": self.entry_city.get(),
            "salary": self.entry_salary.get(),
            "salary_currency": self.cb_currency.get(),
            "experience": self.cb_exp.get(),
            "schedule": self.cb_sched.get(),
            "period": self.cb_period.get(),
            "exclude": self.entry_exclude.get(),
            "enabled_sites": self.get_enabled_sites(),
            "salary_mode": salary_mode,
            "strict_salary": salary_mode == "Строгий",
        }

        self.saved_params[profile_name] = params
        self.save_saved_params()

        self.cb_profiles["values"] = list(self.saved_params.keys())
        self.entry_profile_name.delete(0, tk.END)

        messagebox.showinfo("✓ Успех", f"Профиль '{profile_name}' сохранен!")

    def load_search_params(self, event=None):
        profile_name = self.cb_profiles.get()
        if not profile_name or profile_name not in self.saved_params:
            return

        params = self.saved_params[profile_name]
        self.entry_prof.delete(0, tk.END)
        self.entry_prof.insert(0, params.get("queries", ""))
        self.entry_city.delete(0, tk.END)
        self.entry_city.insert(0, params.get("city", ""))
        self.entry_salary.delete(0, tk.END)
        self.entry_salary.insert(0, params.get("salary", ""))
        self.entry_exclude.delete(0, tk.END)
        self.entry_exclude.insert(0, params.get("exclude", ""))

        enabled_sites = params.get("enabled_sites", [])
        for site_name, var in self.site_vars.items():
            var.set(site_name in enabled_sites)

        self.cb_currency.set(params.get("salary_currency", "BYN"))
        self.cb_exp.set(params.get("experience", "Не имеет значения"))

        schedule = params.get("schedule", "Любой")
        if schedule not in self.cb_sched["values"]:
            schedule = "Любой"
        self.cb_sched.set(schedule)

        period = params.get("period", "За неделю")
        if period in {"За все время", "За месяц"}:
            period = "За 30 дней"
        if period not in self.cb_period["values"]:
            period = "За неделю"
        self.cb_period.set(period)

        salary_mode = params.get("salary_mode")
        if salary_mode not in self.cb_salary_mode["values"]:
            salary_mode = "Строгий" if params.get("strict_salary") else "Мягкий"
        self.cb_salary_mode.set(salary_mode)

    def delete_search_params(self):
        """Удаление профиля"""
        profile_name = self.cb_profiles.get()
        if not profile_name:
            messagebox.showwarning("Внимание", "Выберите профиль!")
            return
        
        if messagebox.askyesno("Подтверждение", f"Удалить профиль '{profile_name}'?"):
            del self.saved_params[profile_name]
            self.save_saved_params()
            
            self.cb_profiles['values'] = list(self.saved_params.keys())
            self.cb_profiles.set('')
            
            messagebox.showinfo("✓ Успех", f"Профиль удален")

    # --- МЕТОДЫ РАБОТЫ С САЙТАМИ ---
    
    def get_enabled_sites(self):
        return [site_name for site_name, var in self.site_vars.items() if var.get()]

    def select_all_sites(self):
        for var in self.site_vars.values():
            var.set(True)

    def deselect_all_sites(self):
        for var in self.site_vars.values():
            var.set(False)

    def select_belarus_sites(self):
        self.deselect_all_sites()
        for site_name in JobScraper.BELARUS_SITES:
            if site_name in self.site_vars:
                self.site_vars[site_name].set(True)

    def select_russia_sites(self):
        self.deselect_all_sites()
        if "HH.ru" in self.site_vars:
            self.site_vars["HH.ru"].set(True)

    # --- МЕТОДЫ ПОИСКА ---
    
    def start_search_thread(self):
        """Validate input and start one background search."""
        if self.search_thread and self.search_thread.is_alive():
            messagebox.showinfo("Поиск", "Поиск уже выполняется.")
            return

        salary_text = self.entry_salary.get().strip()
        if salary_text:
            try:
                salary_value = int(salary_text)
                if salary_value < 0:
                    raise ValueError
            except ValueError:
                messagebox.showwarning("Внимание", "Минимальный доход должен быть целым неотрицательным числом.")
                return

        salary_mode = self.cb_salary_mode.get() or "Мягкий"
        params = {
            "queries": self.entry_prof.get().strip(),
            "city": self.entry_city.get().strip(),
            "exclude": self.entry_exclude.get().strip(),
            "salary": salary_text,
            "salary_currency": self.cb_currency.get(),
            "experience": self.cb_exp.get(),
            "schedule": self.cb_sched.get(),
            "period": self.cb_period.get(),
            "enabled_sites": self.get_enabled_sites(),
            "salary_mode": salary_mode,
            "strict_salary": salary_mode == "Строгий",
        }

        if not params["queries"]:
            messagebox.showwarning("Внимание", "Введите хотя бы одну профессию.")
            return
        if not params["city"]:
            messagebox.showwarning("Внимание", "Введите хотя бы один город.")
            return
        if not params["enabled_sites"]:
            messagebox.showwarning("Внимание", "Выберите хотя бы один источник.")
            return

        self.search_log = []
        self.last_search_error = None
        self.last_search_traceback = None
        self.search_cancel_event.clear()
        try:
            self.current_log_session_dir = self.problem_logger.start_session(params)
            self._append_search_log(f"Логи проблем этой сессии: {self.current_log_session_dir}")
        except OSError as exc:
            self.current_log_session_dir = None
            self.search_log.append(
                f"[{datetime.now():%H:%M:%S}] Не удалось создать папку логов проблем: {exc}"
            )

        self.btn_search.config(state=tk.DISABLED)
        self.btn_stop.config(state=tk.NORMAL)
        self.lbl_status.config(text="⏳ Идет поиск...", fg="blue")
        self.progress["value"] = 0

        for row in self.tree.get_children():
            self.tree.delete(row)
        self.found_data = []
        self.btn_open.config(state=tk.DISABLED)
        self.save_current_settings()

        self.search_thread = threading.Thread(
            target=self.run_search,
            args=(params,),
            daemon=True,
            name="vacancy-search",
        )
        self.search_thread.start()

    def run_search(self, params):
        """Worker-thread search body."""
        try:
            self.found_data = self.scraper.search(
                params,
                progress_callback=self.update_progress,
                cancel_event=self.search_cancel_event,
            )
        except Exception as exc:
            self.last_search_error = f"{type(exc).__name__}: {exc}"
            self.last_search_traceback = traceback.format_exc()
            self._append_search_log(f"Критическая ошибка поиска: {self.last_search_error}")
            self._append_search_log("TRACEBACK: " + self.last_search_traceback.replace("\n", " | "))
            self.found_data = []
        finally:
            try:
                self.root.after(0, self.update_ui_after_search)
            except tk.TclError:
                pass

    def cancel_search(self):
        """Request cooperative cancellation."""
        if self.search_thread and self.search_thread.is_alive():
            self.search_cancel_event.set()
            self.btn_stop.config(state=tk.DISABLED)
            self.lbl_status.config(text="⏹ Останавливаю поиск после текущего запроса...", fg="#b36b00")

    def _append_search_log(self, message):
        """Thread-safe enough in-memory bounded diagnostic log."""
        stamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{stamp}] {message}"
        self.search_log.append(line)
        try:
            self.problem_logger.append(line)
        except Exception:
            pass
        if len(self.search_log) > 1200:
            del self.search_log[:200]

    def _save_search_log_file(self):
        """Финализировать текущую папку диагностической сессии поиска."""
        cancelled = self.search_cancel_event.is_set()
        source_counts = self.scraper.source_counts(self.found_data or [])
        source_errors = self.scraper.source_errors()
        source_stats = self.scraper.get_source_diagnostics()
        status = self.scraper.search_status(
            self.found_data or [],
            cancelled=cancelled,
            fatal_error=self.last_search_error,
        )

        try:
            self.problem_logger.finalize(
                status=status,
                result_count=len(self.found_data or []),
                error=self.last_search_error,
                fatal_traceback=self.last_search_traceback,
                cancelled=cancelled,
                source_counts=source_counts,
                source_errors=source_errors,
                source_stats=source_stats,
                requested_sources=list(self.scraper.requested_sources),
            )
        except Exception:
            # Проблема с диагностикой не должна ломать основной сценарий поиска.
            pass

    def update_progress(self, current, total, message):
        progress_percent = (current / total * 100) if total > 0 else 0
        try:
            self.root.after(0, lambda: self.progress.config(value=progress_percent))
            self.root.after(0, lambda: self.lbl_status.config(text=f"⏳ {message} ({current}/{total})"))
        except tk.TclError:
            pass

    def update_ui_after_search(self):
        cancelled = self.search_cancel_event.is_set()
        status = self.scraper.search_status(
            self.found_data,
            cancelled=cancelled,
            fatal_error=self.last_search_error,
        )
        source_errors = self.scraper.source_errors()
        source_warnings = self.scraper.source_warnings()

        if status == "error":
            if self.last_search_error:
                text = f"❌ Ошибка поиска: {self.last_search_error}"
            elif source_errors:
                text = "❌ Источники недоступны: " + ", ".join(source_errors)
            else:
                text = "❌ Поиск завершился с ошибкой"
            self.lbl_status.config(text=text, fg="red")
            self.progress["value"] = 0
        elif status == "cancelled":
            if self.found_data:
                self.refresh_tree()
                self.lbl_status.config(
                    text=f"⏹ Поиск остановлен. Сохранено результатов: {len(self.found_data)}",
                    fg="#b36b00",
                )
            else:
                self.lbl_status.config(text="⏹ Поиск остановлен", fg="#b36b00")
        elif status == "no_results":
            self.lbl_status.config(text="❌ Ничего не найдено", fg="red")
            self.progress["value"] = 0
        elif status == "partial_success":
            self.refresh_tree()
            self.lbl_status.config(
                text=(
                    f"⚠ Найдено: {len(self.found_data)}. "
                    f"С ошибкой источников: {len(source_errors)} ({', '.join(source_errors)})"
                ),
                fg="#b36b00",
            )
            self.progress["value"] = 100
        else:
            self.refresh_tree()
            if source_warnings:
                warning_sources = ", ".join(source_warnings)
                self.lbl_status.config(
                    text=(
                        f"✓ Найдено: {len(self.found_data)}. "
                        f"Предупреждения: {len(source_warnings)} ({warning_sources})"
                    ),
                    fg="#7a6400",
                )
            else:
                self.lbl_status.config(
                    text=f"✓ Найдено: {len(self.found_data)} уникальных вакансий",
                    fg="green",
                )
            self.progress["value"] = 100

        self.btn_open.config(state=tk.NORMAL if self.found_data else tk.DISABLED)
        self.btn_search.config(state=tk.NORMAL)
        self.btn_stop.config(state=tk.DISABLED)
        self._save_search_log_file()
        self.save_current_settings()

    @staticmethod
    def _display_cell(value):
        """Treeview is single-line: collapse embedded newlines/tabs from websites."""
        return " ".join(str(value or "").replace("\u00a0", " ").replace("\u202f", " ").split())

    def refresh_tree(self):
        for row in self.tree.get_children():
            self.tree.delete(row)
        for item in self.found_data:
            self.tree.insert(
                "",
                tk.END,
                values=(
                    self._display_cell(item.get("position", "")),
                    self._display_cell(item.get("salary", "")),
                    self._display_cell(item.get("company", "")),
                    self._display_cell(item.get("city", "")),
                    self._display_cell(item.get("source", "")),
                    self._display_cell(item.get("date", "")),
                ),
            )

    def refresh_results(self):
        """Обновление результатов"""
        if self.found_data:
            self.refresh_tree()
            messagebox.showinfo("✓", "Результаты обновлены!")

    # --- МЕТОДЫ СОРТИРОВКИ ---
    
    def sort_column(self, col, reverse=False):
        if not self.found_data:
            return

        if self.last_sort_col == col:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_reverse = False
            self.last_sort_col = col
        reverse = self.sort_reverse

        if col == "salary":
            self.found_data.sort(
                key=lambda x: self.scraper.salary_sort_value(x.get("salary", "")),
                reverse=reverse,
            )
        elif col == "date":
            self.found_data.sort(key=lambda x: str(x.get("date", "")), reverse=reverse)
        else:
            self.found_data.sort(
                key=lambda x: str(x.get(col, "")).casefold(),
                reverse=reverse,
            )

        for column in ("position", "salary", "company", "city", "source", "date"):
            text = self.tree.heading(column, "text").replace(" ▲", "").replace(" ▼", "")
            if column == col:
                text += " ▼" if reverse else " ▲"
            self.tree.heading(column, text=text)
        self.refresh_tree()

    # --- МЕТОДЫ КОНТЕКСТНОГО МЕНЮ ---
    
    def show_context_menu(self, event):
        """Показ контекстного меню для таблицы"""
        item = self.tree.identify_row(event.y)
        if item:
            self.tree.selection_set(item)
            self.context_menu.post(event.x_root, event.y_root)

    def show_text_context_menu(self, event):
        """Показ контекстного меню для текста"""
        try:
            self.text_context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.text_context_menu.grab_release()

    def on_key_press(self, event):
        """Обработка клавиш"""
        if event.keycode in [67, 1057] and event.state & 0x4:  # Ctrl+C
            self.copy_text()
            return "break"
        elif event.keycode in [86, 1052] and event.state & 0x4:  # Ctrl+V
            self.paste_text()
            return "break"

    def copy_text(self):
        """Копирование текста"""
        try:
            widget = self.root.focus_get()
            if hasattr(widget, 'get'):
                try:
                    selected_text = widget.selection_get()
                    self.root.clipboard_clear()
                    self.root.clipboard_append(selected_text)
                except tk.TclError:
                    try:
                        full_text = widget.get()
                        self.root.clipboard_clear()
                        self.root.clipboard_append(full_text)
                    except (tk.TclError, AttributeError, TypeError):
                        pass
        except (tk.TclError, AttributeError, TypeError):
            pass

    def paste_text(self):
        """Вставка текста"""
        try:
            widget = self.root.focus_get()
            if hasattr(widget, 'insert'):
                clipboard_text = self.root.clipboard_get()
                if hasattr(widget, 'delete'):
                    try:
                        start = widget.index(tk.SEL_FIRST)
                        end = widget.index(tk.SEL_LAST)
                        widget.delete(start, end)
                    except tk.TclError:
                        pass
                widget.insert(tk.INSERT, clipboard_text)
        except (tk.TclError, AttributeError, TypeError):
            pass

    def select_all_text(self):
        """Выделение всего текста"""
        try:
            widget = self.root.focus_get()
            if hasattr(widget, 'select_range'):
                widget.select_range(0, tk.END)
        except (tk.TclError, AttributeError, TypeError):
            pass

    def copy_selected_to_clipboard(self, event=None):
        """Копирование выбранных вакансий"""
        selection = self.tree.selection()
        if not selection:
            return
        
        text_lines = []
        for item_id in selection:
            vals = self.tree.item(item_id)['values']
            children = self.tree.get_children()
            idx = children.index(item_id)
            
            vacancy = self.found_data[idx]
            line = f"{vacancy['position']} | {vacancy['salary']} | {vacancy['company']} | {vacancy['link']}"
            text_lines.append(line)
        
        self.root.clipboard_clear()
        self.root.clipboard_append('\n'.join(text_lines))
        self.lbl_status.config(text=f"✓ Скопировано {len(text_lines)} вакансий", fg="green")

    def copy_link(self):
        """Копирование ссылки"""
        selection = self.tree.selection()
        if not selection:
            return
        
        item_id = selection[0]
        children = self.tree.get_children()
        idx = children.index(item_id)
        
        link = self.found_data[idx].get('link')
        if link:
            self.root.clipboard_clear()
            self.root.clipboard_append(link)
            self.lbl_status.config(text="✓ Ссылка скопирована", fg="green")

    # --- МЕТОДЫ РАБОТЫ С ВАКАНСИЯМИ ---
    
    def open_selected_vacancy(self, event=None):
        """Open selected vacancies immediately without blocking the GUI.

        GSZ detail links produced by the search parser are already browser-ready.
        No HTTP preflight is performed here: on the real portal a Python request
        can time out while the exact same URL opens normally in Chrome.
        """
        if self._vacancy_open_in_progress:
            self.lbl_status.config(
                text="⏳ Вакансия уже открывается — подождите.",
                fg="#555555",
            )
            return

        selection = self.tree.selection()
        if not selection:
            messagebox.showwarning("Внимание", "Выберите вакансию!")
            return

        children = list(self.tree.get_children())
        selected = []
        for item_id in selection:
            try:
                idx = children.index(item_id)
            except ValueError:
                continue
            if 0 <= idx < len(self.found_data):
                selected.append((idx, dict(self.found_data[idx])))

        if not selected:
            return

        self._vacancy_open_in_progress = True
        self.btn_open.config(state=tk.DISABLED)
        self.lbl_status.config(text="⏳ Открываю вакансию...", fg="#555555")
        threading.Thread(
            target=self._open_vacancies_worker,
            args=(selected,),
            daemon=True,
            name="vacancy-link-opener",
        ).start()

    def _open_vacancies_worker(self, selected):
        opened = 0
        failed = 0
        direct_gsz = 0

        for order, (idx, vacancy) in enumerate(selected):
            original_link = str(vacancy.get("link") or "").strip()
            source = str(vacancy.get("source") or "")
            if not original_link:
                failed += 1
                continue

            link = self.scraper.prepare_vacancy_link_for_browser(original_link, source)
            event = None
            if source == "GSZ.gov.by":
                direct_gsz += 1
                event = {
                    "source": source,
                    "position": vacancy.get("position"),
                    "city": vacancy.get("city"),
                    "original": original_link,
                    "prepared": link,
                    "search_marker_preserved": "source=search" in link,
                    "preflight_skipped": True,
                    "action": "open_direct_detail",
                    "reason": (
                        "GSZ search-result detail URL is opened directly; "
                        "requests preflight intentionally skipped to avoid portal timeout latency"
                    ),
                }
                try:
                    if link and link != original_link:
                        self.found_data[idx]["link"] = link
                except (IndexError, TypeError):
                    pass

            started = time.perf_counter()
            try:
                browser_result = webbrowser.open(link)
                elapsed_ms = round((time.perf_counter() - started) * 1000, 1)
                opened += 1
                if event is not None:
                    event["browser_open_result"] = bool(browser_result)
                    event["open_call_ms"] = elapsed_ms
            except Exception as exc:
                elapsed_ms = round((time.perf_counter() - started) * 1000, 1)
                failed += 1
                if event is None:
                    event = {
                        "source": source,
                        "position": vacancy.get("position"),
                        "original": original_link,
                    }
                event.update({
                    "action": "browser_open_failed",
                    "error": str(exc),
                    "open_call_ms": elapsed_ms,
                })

            if event is not None:
                try:
                    self.problem_logger.append_link_check(event)
                    self.problem_logger.append(
                        f"[{datetime.now():%H:%M:%S}] Vacancy link open: "
                        + json.dumps(event, ensure_ascii=False)
                    )
                except Exception:
                    pass

            # Avoid bursting many tabs when multiple rows are selected, while a
            # single vacancy opens without the old artificial 0.3 s delay.
            if order + 1 < len(selected):
                time.sleep(0.1)

        def finish():
            try:
                self._vacancy_open_in_progress = False
                self.btn_open.config(state=tk.NORMAL if self.found_data else tk.DISABLED)
                if failed:
                    self.lbl_status.config(
                        text=f"⚠ Открыто: {opened}. Ошибок открытия: {failed}.",
                        fg="#b36b00",
                    )
                elif direct_gsz:
                    self.lbl_status.config(
                        text=f"✓ Открыто вакансий: {opened}. GSZ открыт напрямую без проверки ссылки.",
                        fg="green",
                    )
                else:
                    self.lbl_status.config(text=f"✓ Открыто вакансий: {opened}", fg="green")
            except tk.TclError:
                pass

        try:
            self.root.after(0, finish)
        except tk.TclError:
            self._vacancy_open_in_progress = False

    def on_double_click(self, event):
        """Двойной клик"""
        self.open_selected_vacancy()

    def show_vacancy_info(self):
        """Показ информации о вакансии"""
        selection = self.tree.selection()
        if not selection:
            return
        
        item_id = selection[0]
        children = self.tree.get_children()
        idx = children.index(item_id)
        
        vacancy = self.found_data[idx]
        
        info = f"""
Должность: {vacancy['position']}
Зарплата: {vacancy['salary']}
Компания: {vacancy['company']}
Город: {vacancy['city']}
Источник: {vacancy['source']}
Дата: {vacancy['date']}
Ссылка: {vacancy['link']}
        """
        
        messagebox.showinfo("Информация о вакансии", info.strip())

    def show_statistics(self):
        if not self.found_data:
            messagebox.showwarning("Внимание", "Нет данных для статистики!")
            return

        total = len(self.found_data)
        sources = {}
        cities = {}
        for item in self.found_data:
            sources[item.get("source", "Н/Д")] = sources.get(item.get("source", "Н/Д"), 0) + 1
            cities[item.get("city", "Н/Д")] = cities.get(item.get("city", "Н/Д"), 0) + 1

        with_salary = sum(1 for item in self.found_data if self.scraper.parse_salary(item.get("salary", "")))
        salaries = [
            self.scraper.salary_sort_value(item.get("salary", ""))
            for item in self.found_data
        ]
        salaries = [value for value in salaries if value > 0]
        avg_salary = sum(salaries) / len(salaries) if salaries else 0
        median_salary = median(salaries) if salaries else 0

        rates_note = ""
        if self.scraper.currency_rates_updated_at:
            rates_note = f"\nКурсы НБРБ обновлены: {self.scraper.currency_rates_updated_at:%d.%m.%Y %H:%M}"
        elif salaries:
            rates_note = "\nДля части иностранных зарплат пересчёт может быть недоступен."

        stats_text = f"""📊 СТАТИСТИКА ПОИСКА

    Всего уникальных вакансий: {total}
    С указанной зарплатой: {with_salary} ({with_salary / total * 100:.1f}%)
    Средняя оценочная зарплата: {avg_salary:.0f} BYN
    Медианная оценочная зарплата: {median_salary:.0f} BYN{rates_note}

    📍 По городам:
    {self._format_dict(cities, limit=10)}

    🌐 По источникам:
    {self._format_dict(sources)}
    """

        stats_window = tk.Toplevel(self.root)
        stats_window.title("📊 Статистика")
        stats_window.geometry("560x620")

        text_widget = tk.Text(stats_window, wrap=tk.WORD, font=("Segoe UI", 10))
        text_widget.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        text_widget.insert("1.0", stats_text)
        text_widget.config(state=tk.DISABLED)
        ttk.Button(stats_window, text="Закрыть", command=stats_window.destroy).pack(pady=10)

    def show_search_log(self):
        """Показ последнего диагностического лога поиска."""
        if not self.search_log:
            messagebox.showinfo("Лог поиска", "Лог пуст. Выполните поиск.")
            return

        log_window = tk.Toplevel(self.root)
        log_window.title("📝 Лог поиска")
        log_window.geometry("850x620")

        text_frame = tk.Frame(log_window)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        text_widget = tk.Text(text_frame, wrap=tk.WORD, font=("Consolas", 9))
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=text_widget.yview)
        text_widget.configure(yscrollcommand=scrollbar.set)
        text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        text_widget.insert("1.0", "\n".join(self.search_log))
        text_widget.config(state=tk.DISABLED)

        btn_frame = tk.Frame(log_window)
        btn_frame.pack(pady=(0, 10))
        ttk.Button(
            btn_frame,
            text="Копировать",
            command=lambda: self._copy_log_to_clipboard(text_widget),
        ).pack(side=tk.LEFT, padx=5)
        ttk.Button(
            btn_frame,
            text="Открыть папку логов",
            command=self.open_logs_folder,
        ).pack(side=tk.LEFT, padx=5)
        ttk.Button(
            btn_frame,
            text="Очистить текущий",
            command=lambda: self._clear_log(log_window),
        ).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Закрыть", command=log_window.destroy).pack(side=tk.LEFT, padx=5)

    def open_logs_folder(self):
        try:
            LOG_DIR.mkdir(parents=True, exist_ok=True)
            if os.name == "nt":
                os.startfile(str(LOG_DIR))
            else:
                webbrowser.open(LOG_DIR.as_uri())
        except OSError as exc:
            messagebox.showerror("Ошибка", f"Не удалось открыть папку логов:\\n{exc}")

    def _copy_log_to_clipboard(self, text_widget):
        """Копирование лога в буфер обмена"""
        log_text = text_widget.get('1.0', tk.END)
        self.root.clipboard_clear()
        self.root.clipboard_append(log_text)
        messagebox.showinfo("✓", "Лог скопирован в буфер обмена")

    def _clear_log(self, window):
        """Очистка лога"""
        if messagebox.askyesno("Подтверждение", "Очистить лог поиска?"):
            self.search_log = []
            window.destroy()
            messagebox.showinfo("✓", "Лог очищен")

    def _format_dict(self, d, limit=None):
        items = sorted(d.items(), key=lambda x: x[1], reverse=True)
        if limit:
            items = items[:limit]
        return "\n".join(f"  • {key}: {value}" for key, value in items)

    # --- МЕТОДЫ СОХРАНЕНИЯ ---
    
    def save_to_excel(self):
        """Экспорт результатов в аккуратный XLSX с фильтрами и сводкой."""
        if not self.found_data:
            messagebox.showwarning("Внимание", "Нет данных для сохранения!")
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"Vacancies_{timestamp}.xlsx",
        )
        if not file_path:
            return

        def safe_excel_text(value):
            value = "" if value is None else str(value)
            if value.startswith(("=", "+", "-", "@")):
                return "'" + value
            return value

        try:
            export_data = []
            for item in self.found_data:
                export_data.append({
                    "Должность": safe_excel_text(item.get("position")),
                    "Зарплата": safe_excel_text(item.get("salary")),
                    "Компания": safe_excel_text(item.get("company")),
                    "Город": safe_excel_text(item.get("city")),
                    "Источник": safe_excel_text(item.get("source")),
                    "Дата": safe_excel_text(item.get("date")),
                    "Ссылка": str(item.get("link") or ""),
                })

            df = pd.DataFrame(export_data)
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="Вакансии", index=False)
                worksheet = writer.sheets["Вакансии"]
                worksheet.freeze_panes = "A2"
                worksheet.sheet_view.showGridLines = False

                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                for idx, column_name in enumerate(df.columns, start=1):
                    values = [str(column_name)] + [str(v) for v in df[column_name].fillna("").tolist()]
                    max_len = min(max(len(v) for v in values) + 2, 70)
                    worksheet.column_dimensions[get_column_letter(idx)].width = max(12, max_len)

                for row in worksheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)

                link_col_idx = list(df.columns).index("Ссылка") + 1
                for row_idx in range(2, len(df) + 2):
                    cell = worksheet.cell(row=row_idx, column=link_col_idx)
                    link_url = str(cell.value or "")
                    if link_url.startswith(("http://", "https://")):
                        cell.hyperlink = link_url
                        cell.font = Font(color="0563C1", underline="single")

                if len(df) > 0:
                    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
                    table = Table(displayName="VacanciesTable", ref=ref)
                    table.tableStyleInfo = TableStyleInfo(
                        name="TableStyleMedium2",
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=True,
                        showColumnStripes=False,
                    )
                    worksheet.add_table(table)

                summary_rows = [
                    ["Параметр", "Значение"],
                    ["Дата экспорта", datetime.now().strftime("%d.%m.%Y %H:%M:%S")],
                    ["Уникальных вакансий", len(self.found_data)],
                    ["С указанной зарплатой", sum(1 for i in self.found_data if self.scraper.parse_salary(i.get("salary", "")))],
                ]
                source_counts = {}
                for item in self.found_data:
                    source = item.get("source", "Н/Д")
                    source_counts[source] = source_counts.get(source, 0) + 1
                for source, count in sorted(source_counts.items(), key=lambda pair: pair[1], reverse=True):
                    summary_rows.append([f"Источник: {source}", count])

                summary = writer.book.create_sheet("Сводка")
                for row in summary_rows:
                    summary.append(row)
                for cell in summary[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                summary.column_dimensions["A"].width = 35
                summary.column_dimensions["B"].width = 25
                summary.freeze_panes = "A2"

            messagebox.showinfo("✓ Успех", f"Файл сохранён.\nЭкспортировано: {len(export_data)} вакансий")
            self.lbl_status.config(text=f"✓ Сохранено: {file_path}", fg="green")
        except Exception as exc:
            messagebox.showerror("Ошибка", f"Не удалось сохранить Excel:\n{exc}")


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
